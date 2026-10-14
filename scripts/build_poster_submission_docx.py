from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "submission" / "word"
OUT = OUT_DIR / "笑面虎的邀请函-投稿材料.docx"
POSTER = ROOT / "submission" / "poster" / "笑面虎的邀请函-主题海报-300dpi.jpg"

BODY_FONT = "Microsoft YaHei"
ACCENT = "4F315F"
LIGHT_ACCENT = "F2EDF5"
BORDER = "D9D9D9"
GOLD = "9A7412"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, bold=False, color="000000", size=10, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, rows, widths=None):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        if widths:
            cells[0].width = Inches(widths[0])
            cells[1].width = Inches(widths[1])
        set_cell_text(cells[0], label, bold=True, color=ACCENT, size=9.5)
        set_cell_shading(cells[0], LIGHT_ACCENT)
        set_cell_text(cells[1], value, size=9.5)
        if idx % 2 == 1:
            set_cell_shading(cells[1], "FAFAFA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True, color="000000")
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.first_line_indent = Inches(0.24)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=10.5, bold=True)
        run = p.add_run(text[len(bold_lead):])
        set_run_font(run, size=10.5)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=GOLD)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8, color="777777")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8, color="777777")


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    body = styles["Body Text"]
    body.font.name = BODY_FONT
    body._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    body.font.size = Pt(10.5)
    for level in (1, 2, 3):
        style = styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    doc.core_properties.title = "笑面虎的邀请函投稿材料"
    doc.core_properties.subject = "校园反诈主题 AI 海报投稿材料"
    doc.core_properties.author = "作者待填写"
    doc.core_properties.comments = "用于校园反诈主题 AI 作品大赛投稿"

    # Use a plain paragraph instead of Word's built-in Title style: the
    # localized template can add an unwanted blue border below the title.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("校园反诈主题 AI 海报投稿材料")
    set_run_font(title_run, size=22, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(13)
    sub_run = subtitle.add_run("笑面虎的邀请函")
    set_run_font(sub_run, size=16, bold=True, color=ACCENT)

    intro = doc.add_paragraph(style="Body Text")
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Inches(0.25)
    intro.paragraph_format.right_indent = Inches(0.25)
    intro.paragraph_format.space_after = Pt(12)
    intro.paragraph_format.line_spacing = 1.3
    intro_run = intro.add_run("本材料用于“智防诈骗·护航青春”校园反诈主题 AI 作品大赛海报投稿。作者信息中的方括号内容需要在发送邮件前替换。")
    set_run_font(intro_run, size=10.5, color="555555")

    if POSTER.exists():
        image_p = doc.add_paragraph()
        image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_p.paragraph_format.space_after = Pt(3)
        image_p.add_run().add_picture(str(POSTER), width=Inches(2.32))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(14)
        cap_run = caption.add_run("投稿海报预览")
        set_run_font(cap_run, size=8.5, color="777777")

    add_heading(doc, "作品信息", 1)
    add_table(doc, [
        ("作品名称", "《笑面虎的邀请函》"),
        ("副标题", "《8.8元到账之后》"),
        ("作品类别", "AI创意海报、插画"),
        ("核心标语", "第一笔返款先给甜头，连环收割随后登场。"),
        ("投稿文件", "笑面虎的邀请函-主题海报-300dpi.jpg"),
        ("文件规格", "3072×4608像素，2:3竖版，JPG；另附同尺寸 PNG 备用"),
        ("分辨率", "300dpi"),
    ], widths=(1.25, 5.55))

    add_heading(doc, "作者信息", 1)
    add_table(doc, [
        ("作者姓名或团队名称", "[填写]"),
        ("所属院系或单位", "[填写]"),
        ("班级", "[填写]"),
        ("联系电话", "[填写]"),
        ("联系邮箱", "[填写]"),
    ], widths=(1.65, 5.15))

    add_heading(doc, "投稿版创作说明", 1)
    add_quote(doc, "《笑面虎的邀请函》以兼职返款为切口：操盘者用红线牵引金币、手机与账户锁，学生走向看似通往远方的门。风景象征诱人的高薪远方，藤蔓暗示控制与失去自由。第一笔返款先给甜头，连环收割随后登场。")
    add_body(doc, "上面这段文字共92字，符合通知中“100字以内创作说明”的要求。")

    add_heading(doc, "设计阐释", 1)
    add_body(doc, "《笑面虎的邀请函》把一张看似诱人的机会海报转化为诈骗机制的可视化叙事。标题中的“笑面虎”指向诈骗话术的伪装性；“邀请函”对应画面中发光的门，它看起来像通往高薪兼职、海外机会或快速赚钱的入口，实际却把人引向逐步失去判断权的诈骗链条。")

    add_heading(doc, "微笑面具与操盘者", 2)
    add_body(doc, "面具始终保持微笑，对应诈骗者常用的客服、学长、同学或平台专员身份，危险的伪装性正体现在这里。诈骗开始时，危险往往以“我来帮你”“这单很简单”“已经有人成功了”的语气出现。画面上方伸出的手形成俯视关系，表现诈骗链中信息和权力的不对等。")

    add_heading(doc, "发光的门与远方风景", 2)
    add_body(doc, "发光的门具有出口和机会的双重意味。门后的风景可以被观众理解为高薪出境、海外工作、快速赚钱或“最后一步就能提现”的远方承诺。作品不把画面指认为某个具体地点，也不声称这是缅北实景；“缅北”等境外园区只是观众可能联想到的现实风险。作品真正要表达的是：看似通往自由的远方，可能通往封闭、失控和被控制的诈骗空间。")

    add_heading(doc, "门边藤蔓与行动受限", 2)
    add_body(doc, "门边的枝状景观在设计阐释中被理解为藤蔓。它们远看像自然风景，靠近后却与红色牵引线交织，象征骗子用环境包装、群体氛围和连续任务慢慢限制人的行动。越向门内走，退路越少；越相信“再完成一步就能提现”，越难停下来。")

    add_heading(doc, "红色牵引线", 2)
    add_body(doc, "红线连接金币、手机、任务凭证和账户锁，代表从8.8元返款到拉新、垫资、假到账、解冻和屏幕共享的控制链。线条起初像装饰，随后变成束缚，表达“你以为自己在做选择，实际上每一步都被话术提前安排”。")

    add_heading(doc, "金币 手机与账户锁", 2)
    add_body(doc, "金色金币和礼物代表第一笔返款带来的奖励感。它看起来像平台诚意，实际承担的是换取信任的诱饵功能。手机是群聊、链接、客服和任务平台的入口；红色账户锁则对应操作错误、账户冻结、缴费解冻和屏幕共享等后续话术。骗子先让人相信页面上的数字，再利用页面异常继续索要钱和权限。")

    add_heading(doc, "走向门的学生与碎片屏幕", 2)
    add_body(doc, "学生采用背影，方便观众把自己代入画面。他没有奔跑，也没有明显恐惧，说明受骗往往从一次普通点击开始：先试一下、反正只有8.8元、同学也做过、再做一单就收手。周围的透明屏幕和碎片化界面象征群聊晒单、转账截图、平台余额和伪造凭证，呈现“证据很多，真正可核验的来源很少”。")

    add_heading(doc, "色彩与构图", 2)
    add_body(doc, "深紫和靛蓝代表数字空间与看似专业的技术包装；金色代表返款和奖励；红色代表牵引、警报、冻结和失去控制；冷白色的门光代表被包装成机会的诱惑。操盘者位于上方，学生位于下方，红线从面具引向金币、手机、账户锁和学生，构成“诱饵—点击—控制”的阅读路径。")

    add_heading(doc, "反诈教育指向", 2)
    add_body(doc, "作品把重点放在骗局如何利用熟人信任、小额返款、群体晒单、时间压力和沉没成本；受害者始终是普通学生，一次普通的点击足以入局。作品希望观众在看到“到账”时先停下来，独立核验来源，不因为同学转发或页面显示数字就继续付款。")
    add_quote(doc, "第一笔返款先给甜头，连环收割随后登场。")

    add_heading(doc, "AI创作及版权说明", 1)
    add_body(doc, "画面由生成式图像模型辅助完成，作者负责主题策划、视觉隐喻设定、构图筛选、色彩方向、文件整理和投稿版输出。人物和场景均为虚构设定，不对应真实个人。作品不使用清华大学校徽、官方印章或校名标准组合，也不宣称为学校官方服务。")

    add_heading(doc, "投稿邮件信息", 1)
    add_table(doc, [
        ("投稿邮箱", "bwbpcs@tsinghua.edu.cn"),
        ("邮件主题", "反诈AI大赛+姓名或团队名称+所在院系+AI创意海报"),
        ("建议附件", "海报 JPG、作者信息、创作说明；可附 AI 创作及版权说明"),
    ], widths=(1.25, 5.55))
    add_body(doc, "老师您好：现提交“智防诈骗·护航青春”校园反诈主题 AI 作品大赛参赛作品。作品名称：《笑面虎的邀请函》，副标题：《8.8元到账之后》，作品类别：AI创意海报、插画。作者、院系、班级和联系方式详见附件，敬请查收。谢谢！")

    add_heading(doc, "提交前检查", 1)
    for item in [
        "把作者姓名或团队名称、所属院系或单位、班级、联系电话和邮箱填入作者信息。",
        "确认附件中的海报文件为笑面虎的邀请函-主题海报-300dpi.jpg。",
        "确认压缩包名称符合通知格式：反诈AI大赛+姓名或团队名称+所在院系+AI创意海报。",
        "确认邮件正文没有把作品写成清华大学官方服务，也没有把画面指认为某个具体境外地点。",
    ]:
        add_bullet(doc, item)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
